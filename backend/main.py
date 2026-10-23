import io
import os
import re
import secrets
from typing import Any

import fitz
import pytesseract
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageOps
from pydantic import BaseModel, Field
from supabase import create_client
from starlette.responses import StreamingResponse

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 5
ALLOWED_TYPES = {"application/pdf", "image/jpeg", "image/png"}
SUPABASE_URL = os.getenv("SUPABASE_URL", "").rstrip("/")
SUPABASE_ANON_KEY = os.getenv("SUPABASE_ANON_KEY", "")
ALLOWED_ORIGINS = [x.strip() for x in os.getenv("ALLOWED_ORIGINS", "http://localhost:8788,http://127.0.0.1:8788").split(",") if x.strip()]

app = FastAPI(title="ConfiCel Conversion API", version="0.5.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Authorization", "Content-Type"],
)

class ExportRequest(BaseModel):
    mode: str = Field(pattern="^(invoice|cheque)$")
    fields: dict[str, Any]
    rows: list[dict[str, Any]] = Field(default_factory=list, max_length=500)

def get_verified_user(authorization: str | None = Header(default=None)) -> dict:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(status_code=401, detail="Sign in to use ConfiCel conversion.")
    if not SUPABASE_URL or not SUPABASE_ANON_KEY:
        raise HTTPException(status_code=503, detail="Authentication is not configured on the API.")
    token = authorization.split(" ", 1)[1].strip()
    try:
        client = create_client(SUPABASE_URL, SUPABASE_ANON_KEY)
        response = client.auth.get_user(token)
        user = getattr(response, "user", None)
        if not user or not getattr(user, "id", None):
            raise ValueError("No verified user")
        return {"id": user.id, "email": getattr(user, "email", None)}
    except Exception:
        raise HTTPException(status_code=401, detail="Your session is invalid or expired. Please sign in again.")

def _ocr_image(image: Image.Image) -> str:
    image = ImageOps.exif_transpose(image).convert("RGB")
    return pytesseract.image_to_string(image, config="--psm 6")[:100_000]

def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise HTTPException(status_code=400, detail="This PDF could not be opened. Check the file and try again.")
    if doc.page_count < 1 or doc.page_count > MAX_PAGES:
        doc.close()
        raise HTTPException(status_code=413, detail=f"PDFs must contain between 1 and {MAX_PAGES} pages.")
    pieces = []
    try:
        for page in doc:
            text = page.get_text("text").strip()
            if len(text) < 20:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                text = _ocr_image(image)
            pieces.append(text)
    finally:
        doc.close()
    return "\n\n".join(pieces).strip(), len(pieces)

def _invoice_draft(text: str) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    date_match = re.search(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", text)
    money = re.findall(r"(?:₹|INR\s*)?\s*([0-9]{1,3}(?:,[0-9]{3})+(?:\.\d{2})?|[0-9]+\.\d{2})", text, re.I)
    return {
        "document_type": "invoice",
        "supplier_or_header": lines[0][:200] if lines else "",
        "invoice_number": "",
        "invoice_date": date_match.group(0) if date_match else "",
        "currency": "INR",
        "total_candidates": money[-5:],
        "line_items": [],
        "raw_text": text,
        "review_required": True,
        "notice": "Draft extraction only. Verify every field and add/correct line items before export."
    }

def _cheque_draft(text: str) -> dict:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    date_match = re.search(r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}-\d{2}-\d{2})\b", text)
    return {
        "document_type": "cheque",
        "bank_or_header": lines[0][:200] if lines else "",
        "cheque_number": "",
        "date_candidate": date_match.group(0) if date_match else "",
        "payee": "",
        "amount_numeric": "",
        "amount_words": "",
        "micr_or_routing_text": "",
        "raw_text": text,
        "review_required": True,
        "notice": "Cheque fields are not reliably inferred yet. Verify all details manually; do not use this as payment validation."
    }

@app.get("/health")
def health():
    return {"ok": True, "service": "conficel-conversion-api", "version": "0.5.0"}

@app.post("/api/convert")
async def convert(file: UploadFile = File(...), mode: str = "invoice", user: dict = Depends(get_verified_user)):
    if mode not in {"invoice", "cheque"}:
        raise HTTPException(status_code=400, detail="Mode must be invoice or cheque.")
    if file.content_type not in ALLOWED_TYPES:
        raise HTTPException(status_code=415, detail="Upload a PDF, JPG, or PNG file.")
    data = await file.read(MAX_BYTES + 1)
    if not data:
        raise HTTPException(status_code=400, detail="The selected file is empty.")
    if len(data) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="Files must be 10 MB or smaller.")
    # Verify actual file signature instead of trusting the browser-provided MIME type.
    if file.content_type == "application/pdf" and not data.startswith(b"%PDF-"):
        raise HTTPException(status_code=400, detail="The file content does not look like a valid PDF.")
    if file.content_type == "image/png" and not data.startswith(b"\x89PNG\r\n\x1a\n"):
        raise HTTPException(status_code=400, detail="The file content does not look like a valid PNG.")
    if file.content_type == "image/jpeg" and not data.startswith(b"\xff\xd8\xff"):
        raise HTTPException(status_code=400, detail="The file content does not look like a valid JPG.")
    try:
        if file.content_type == "application/pdf":
            text, pages = _extract_pdf(data)
        else:
            image = Image.open(io.BytesIO(data))
            if image.width * image.height > 40_000_000:
                raise HTTPException(status_code=413, detail="Image dimensions are too large.")
            text, pages = _ocr_image(image), 1
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read this file. Try a clearer scan or image.")
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text found. Try a clearer scan.")
    draft = _cheque_draft(text) if mode == "cheque" else _invoice_draft(text)
    # No upload is persisted; no credit is deducted in this foundation checkpoint.
    return {"conversion_id": secrets.token_urlsafe(12), "mode": mode, "pages_processed": pages, "draft": draft, "credits_charged": 0, "status": "draft_review", "user_id": user["id"]}

@app.post("/api/export.xlsx")
def export_xlsx(payload: ExportRequest, user: dict = Depends(get_verified_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cheque review" if payload.mode == "cheque" else "Invoice review"
    headers = list(payload.fields.keys())
    ws.append(["Field", "Value"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="174078")
    for key, value in payload.fields.items():
        ws.append([str(key), str(value) if value is not None else ""])
    if payload.rows:
        detail = wb.create_sheet("Line items")
        row_headers = list(dict.fromkeys(k for row in payload.rows for k in row.keys()))
        detail.append(row_headers)
        for cell in detail[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174078")
        for row in payload.rows:
            detail.append([row.get(k, "") for k in row_headers])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in col) + 2, 12), 42)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "conficel-cheque-review.xlsx" if payload.mode == "cheque" else "conficel-invoice-review.xlsx"
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"})

